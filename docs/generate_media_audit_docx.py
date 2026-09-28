"""Generate MEDIA_EVALUATION_TECHNICAL_AUDIT.docx for easy Word use."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parent / "MEDIA_EVALUATION_TECHNICAL_AUDIT.docx"


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18)


def add_h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def add_h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def add_p(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
    doc.add_paragraph()


def add_flow(doc: Document, lines: list[str]) -> None:
    for line in lines:
        doc.add_paragraph(line)


def build() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    add_title(doc, "ACCESS VisionCheck")
    add_title(doc, "Media Evaluation & AI Detection — Technical Audit")
    add_p(doc, "Project: access_mobile")
    add_p(doc, "Date: September 25, 2026")
    add_p(doc, "Purpose: Evidence-based audit of what the codebase actually implements.")
    doc.add_paragraph()

    add_h1(doc, "Executive Summary")
    add_bullets(
        doc,
        [
            "Two separate EfficientNet-B0 models (PyTorch/torchvision), not one dual-output model.",
            "Quality model: 7 technical signals (sharpness, blur, noise, etc.).",
            "AI model: one sigmoid probability (AI-like vs natural).",
            "AI detection is multi-signal fusion (filename, EXIF, OpenCV heuristics, neural ~28%).",
            "Gemini is narrative feedback only — it does not score quality or detect AI.",
            "Training uses synthetic images with frozen ImageNet backbone; no real labeled dataset in repo.",
        ],
    )

    add_h1(doc, "1. Relevant Files")
    add_h2(doc, "Backend — Core Pipeline")
    add_table(
        doc,
        ["File", "Role"],
        [
            ["access_backend/app/api/routes/media.py", "POST /media/evaluate, POST /media/upload"],
            ["access_backend/app/services/media_pipeline.py", "Orchestrates evaluation, DB, XP, alerts"],
            ["access_backend/app/services/media_evaluator.py", "EfficientNet + FFmpeg + AI detection"],
            ["access_backend/app/services/efficientnet_b0_scorer.py", "Quality: EfficientNet-B0 + 7-output head"],
            ["access_backend/app/services/ai_authenticity_model.py", "AI: EfficientNet-B0 + 1-output sigmoid"],
            ["access_backend/app/services/ai_authenticity_detector.py", "Multi-signal AI verdict fusion"],
            ["access_backend/app/services/ffmpeg_frames.py", "Video frames: FFmpeg or OpenCV fallback"],
            ["access_backend/app/services/opencv_quality.py", "OpenCV fallback quality metrics"],
            ["access_backend/app/services/gemini_service.py", "Gemini narrative only"],
        ],
    )

    add_h2(doc, "Models & Training")
    add_table(
        doc,
        ["File", "Role"],
        [
            ["access_backend/ml/efficientnet_b0_heads.pt", "Quality head weights"],
            ["access_backend/ml/ai_authenticity_head.pt", "AI authenticity head weights"],
            ["access_backend/scripts/train_efficientnet_heads.py", "Train quality on synthetic degradations"],
            ["access_backend/scripts/train_ai_authenticity_head.py", "Train AI on synthetic natural vs AI-like"],
        ],
    )

    add_h2(doc, "Flutter — Production Client")
    add_table(
        doc,
        ["File", "Role"],
        [
            ["lib/mobile_app/screens/ai_media_evaluation_screen.dart", "Calls POST /media/evaluate, Gemini, upload"],
            ["lib/shared/api/member_api_service.dart", "evaluateMedia(), fetchGeminiMediaFeedback()"],
            ["lib/shared/services/media_validation_service.dart", "Pre-upload client validation only"],
        ],
    )

    add_h2(doc, "Legacy / Unused")
    add_table(
        doc,
        ["File", "Role"],
        [
            ["lib/shared/services/efficientnet_scoring_service.dart", "Local mirror — not called by eval screen"],
            ["lib/shared/api/media_pipeline.dart", "Simulated pipeline — not used in production"],
        ],
    )

    add_h1(doc, "2. Image Upload → Final Result")
    add_flow(
        doc,
        [
            "User picks image (Flutter)",
            "→ MediaValidationService.validatePhoto() [client pre-check]",
            "→ POST /api/media/evaluate",
            "→ evaluate_bytes() → evaluate_photo()",
            "→ _analyze_image() → efficientnet_b0_scorer.analyze_image_bytes()",
            "→ _photo_criteria() + _quality_warnings()",
            "→ _build_payload() → detect_ai_generated()",
            "→ apply_integrity_evaluation_policy()",
            "→ Optional Gemini narrative",
            "→ POST /api/media/upload + evaluation_metadata",
            "→ run_pipeline() → DB: media_evaluations + ai_detection_results",
        ],
    )

    add_h2(doc, "Image Pipeline Steps")
    add_table(
        doc,
        ["Step", "File", "Function", "Library", "Output"],
        [
            ["1", "media_validation_service.dart", "validatePhoto", "image package", "Accept/reject"],
            ["2", "member_api_service.dart", "evaluateMedia", "HTTP", "JSON from backend"],
            ["3", "media.py", "evaluate_media", "FastAPI", "Validated upload"],
            ["4", "media_evaluator.py", "evaluate_photo", "—", "Metric dict"],
            ["5", "efficientnet_b0_scorer.py", "analyze_image_bytes", "PyTorch/torchvision", "7 metrics 0–1"],
            ["6", "media_evaluator.py", "_photo_criteria", "Rule-based", "7 criteria scores 1–5"],
            ["7", "ai_authenticity_detector.py", "detect_ai_generated", "OpenCV + PyTorch", "ai_detection dict"],
            ["8", "media_pipeline.py", "run_pipeline", "SQLAlchemy", "DB rows + XP"],
        ],
    )

    add_h1(doc, "3. Video Upload → Final Result")
    add_flow(
        doc,
        [
            "User picks video → validateVideo (2 frames client-side)",
            "→ POST /api/media/evaluate → evaluate_video()",
            "→ ffmpeg_frames.extract_frames() [up to 6 JPEG frames]",
            "→ Per frame: _analyze_image() [EfficientNet or OpenCV]",
            "→ _average_metrics() + _flickering() + _motion_stability()",
            "→ detect_ai_generated_from_frames() [≤5 frames, max probability]",
            "→ upload → run_pipeline() → final result",
        ],
    )

    add_h2(doc, "Video Processing Details")
    add_table(
        doc,
        ["Question", "Answer"],
        [
            ["Decoding", "FFmpeg subprocess, or OpenCV VideoCapture fallback"],
            ["Frames extracted", "Yes — default 6 (DEFAULT_SAMPLE_COUNT = 6)"],
            ["Sampling", "Evenly spaced by duration via ffprobe"],
            ["Resize for model", "224×224 RGB, ImageNet normalize"],
            ["Per-frame analysis", "Yes — each frame through _analyze_image()"],
            ["Quality combination", "Arithmetic mean of 7 metrics; width/height from first frame"],
            ["Video criteria count", "9 total (7 averaged + flickering + motion stability)"],
            ["AI on video", "Up to 5 frames; maximum synthetic probability wins"],
            ["FPS / bitrate", "Not implemented"],
        ],
    )

    add_h1(doc, "4. EfficientNet-B0 Investigation")
    add_h2(doc, "Model A — Quality (IQA)")
    add_table(
        doc,
        ["Property", "Value"],
        [
            ["Library", "PyTorch + torchvision (efficientnet_b0)"],
            ["Pretrained", "EfficientNet_B0_Weights.IMAGENET1K_V1"],
            ["Weights file", "access_backend/ml/efficientnet_b0_heads.pt"],
            ["Architecture", "Backbone + Dropout(0.15) + Linear(1280→7) + sigmoid"],
            ["Input", "224×224 RGB"],
            ["Output", "7 values 0–1: sharpness, blur, noise, brightness, contrast, blocking, ringing"],
        ],
    )

    add_h2(doc, "Model B — AI Authenticity")
    add_table(
        doc,
        ["Property", "Value"],
        [
            ["Library", "PyTorch + torchvision"],
            ["Pretrained", "ImageNet-1K V1"],
            ["Weights file", "access_backend/ml/ai_authenticity_head.pt"],
            ["Architecture", "Backbone + Dropout(0.2) + Linear(1280→1) + sigmoid"],
            ["Output", "P(AI-like) in [0, 1]"],
            ["Training labels", "1.0 = AI-like synthetic, 0.0 = natural synthetic"],
        ],
    )

    add_h2(doc, "Training Evidence")
    add_table(
        doc,
        ["Item", "Quality Model", "AI Model"],
        [
            ["Script", "train_efficientnet_heads.py", "train_ai_authenticity_head.py"],
            ["Dataset", "Synthetic degradations", "Synthetic natural vs AI-like"],
            ["Real photos in repo", "No", "No"],
            ["Backbone fine-tuned", "No — frozen", "No — frozen"],
            ["Head trained", "Yes", "Yes"],
            ["Loss", "MSE", "BCE"],
            ["Optimizer", "Adam lr=1e-3", "Adam lr=1e-3"],
            ["Epochs", "4 default", "5 default"],
        ],
    )
    add_p(
        doc,
        "Finding: No evidence EfficientNet-B0 was fine-tuned on real human-vs-AI-generated photos. "
        "Both heads trained on synthetic data with frozen ImageNet backbone.",
        bold=True,
    )

    add_h1(doc, "5. What EfficientNet-B0 Actually Predicts")
    add_bullets(
        doc,
        [
            "Quality model: 7 continuous technical signals → rule-mapped to 1–5 criteria (not High/Low class).",
            "AI model: P(synthetic/AI-like) — one signal in fusion, not the sole verdict.",
            "AI training uses synthetic AI-like patterns (bilateral filter, oversharpen, upscale), not real Midjourney/DALL·E images.",
        ],
    )

    add_h1(doc, "6. AI-Generated Media Detection")
    add_flow(
        doc,
        [
            "Image bytes",
            "→ Filename hints (Midjourney, DALL·E, ChatGPT, etc.)",
            "→ EXIF/metadata analysis",
            "→ OpenCV heuristics (frequency, noise uniformity, texture entropy, quality paradox, skin smoothness)",
            "→ Optional EfficientNet AI head (weight 0.28 in fusion)",
            "→ _fuse_probability() → _verdict_from_probability()",
            "→ authentic | suspicious | ai_generated",
        ],
    )
    add_table(
        doc,
        ["Threshold", "Value", "Verdict"],
        [
            ["AI_GENERATED_THRESHOLD", "0.58", "ai_generated"],
            ["SUSPICIOUS_THRESHOLD", "0.36", "suspicious (if below 0.58)"],
            ["Below 0.36", "—", "authentic"],
        ],
    )
    add_p(doc, "Gemini does NOT detect AI-generated media.")

    add_h1(doc, "7. Media Quality Metrics")
    add_table(
        doc,
        ["Source", "Metrics", "In final score?"],
        [
            ["EfficientNet-B0 (primary)", "Sharpness, blur, noise, brightness, contrast, blocking, ringing", "Yes"],
            ["OpenCV fallback", "Same 7 names via classical CV", "Yes (if EfficientNet unavailable)"],
            ["FFmpeg (video)", "Frame extraction only; duration for sampling", "Indirect"],
            ["Video heuristics", "Flickering, motion stability", "Yes — 2 extra criteria"],
            ["Not implemented", "FPS, bitrate, frame drops, full temporal model", "No"],
        ],
    )

    add_h1(doc, "8. Final Quality Score")
    add_bullets(
        doc,
        [
            "Photo: overall = mean of 7 criterion scores (each 1–5).",
            "Video: overall = mean of 9 criterion scores.",
            "_soft_score(): boosted = metric × 0.5 + 0.42, bucketed to 1–5.",
            "Excellent ≥4.5, Good ≥3.5, Fair ≥2.5, Needs Improvement <2.5.",
            "Legacy DB overall_score = overall / 5 (0–1 scale).",
            "AI detection does not change quality score; integrity policy hides feedback when flagged.",
        ],
    )

    add_h1(doc, "9. Video-Specific Evaluation")
    add_table(
        doc,
        ["Feature", "Status"],
        [
            ["Video resolution", "Implemented — first frame dimensions"],
            ["Per-frame quality", "Implemented — EfficientNet per frame, averaged"],
            ["Flickering", "Implemented"],
            ["Motion stability", "Implemented"],
            ["FPS", "Not implemented"],
            ["Bitrate", "Not implemented"],
            ["Duration scored", "Partial — ffprobe for sampling only"],
            ["Frame drops", "Not implemented"],
            ["Temporal consistency", "Partial — flickering + motion only"],
        ],
    )

    add_h1(doc, "10. Dual-Output Model?")
    add_p(doc, "NOT IMPLEMENTED.", bold=True)
    add_p(doc, "Two separate EfficientNet-B0 models + separate weight files + OpenCV/EXIF heuristics for AI.")

    add_h1(doc, "11. Confidence Scores")
    add_table(
        doc,
        ["Display", "Source"],
        [
            ["X% AI-generated confidence", "ai_probability when label is AI"],
            ["X% Human confidence", "1 − ai_probability when authentic"],
            ["X% Suspicious confidence", "ai_probability when suspicious"],
            ["Quality scores", "Mean of 1–5 buckets — not softmax probability"],
        ],
    )
    add_p(doc, "No probability calibration found in codebase.")

    add_h1(doc, "12. Dataset")
    add_table(
        doc,
        ["Item", "Finding"],
        [
            ["Real dataset in repo", "None"],
            ["Training data", "Synthetic only, generated at runtime"],
            ["Trained weights in repo", "Yes — both .pt files"],
            ["Train/val/test split", "No persistent split; benchmark uses --seed"],
        ],
    )

    add_h1(doc, "13. Data Leakage")
    add_p(doc, "No classical leakage evidence — no real labeled dataset to split. Main risk: synthetic-to-real domain gap.")

    add_h1(doc, "14. Architecture")
    add_h2(doc, "Current")
    add_flow(
        doc,
        [
            "Flutter → POST /media/evaluate → media_evaluator",
            "Photo: EfficientNet IQA + detect_ai_generated",
            "Video: FFmpeg frames → per-frame IQA → average + temporal heuristics + frame AI scan",
            "→ integrity policy → optional Gemini → upload → DB",
        ],
    )
    add_h2(doc, "Ideal (comparison)")
    add_flow(
        doc,
        [
            "Real labeled IQA data + real human/AI dataset",
            "Calibrated probabilities + validated held-out member uploads",
            "True video IQA or certified external API where needed",
        ],
    )

    add_h1(doc, "15. Final Verdict — 12 Questions")
    qa = [
        ("Q1: Using EfficientNet-B0?", "Yes — two separate PyTorch/torchvision instances."),
        ("Q2: What does it do?", "Quality: 7 regressions. AI: one sigmoid as ~28% of fusion."),
        ("Q3: Trained for AI detection?", "Head trained on synthetic data only; backbone frozen."),
        ("Q4: What detects AI?", "Multi-signal fusion — not Gemini, not EfficientNet alone."),
        ("Q5: Image quality?", "EfficientNet IQA primary; OpenCV fallback."),
        ("Q6: Video quality?", "FFmpeg/OpenCV frames → averaged per-frame metrics + 2 heuristics."),
        ("Q7: Video processing?", "6 frames extracted; AI scans up to 5; max prob wins."),
        ("Q8: Dual-output?", "No."),
        ("Q9: Confidence?", "Fused ai_probability; quality is mean of 1–5 scores."),
        ("Q10: Classifications?", "Quality: 1–5 + labels. Integrity: authentic/suspicious/ai_generated."),
        (
            "Q11: Paper inconsistencies?",
            "UI may claim single EfficientNet for both tasks, 70/30 fusion, creative criteria — code differs.",
        ),
        (
            "Q12: Verify before changes?",
            "Live inference tests; weight loading; FFmpeg on server; DB field semantics; synthetic training vs thesis claims.",
        ),
    ]
    for q, a in qa:
        add_p(doc, q, bold=True)
        add_p(doc, a)

    add_h1(doc, "Bottom Line")
    add_p(
        doc,
        "Production: backend pipeline with two EfficientNet-B0 heads, FFmpeg video sampling, "
        "OpenCV/heuristic AI fusion, Gemini for coaching text only. "
        "NOT a dual-output model. AI detection is NOT EfficientNet-only.",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Created: {path}")
